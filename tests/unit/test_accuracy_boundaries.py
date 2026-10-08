"""Regression cases from the 2026-09-22 audit. All fixtures are synthetic."""
from datetime import date, datetime
from dataclasses import replace
from decimal import Decimal as D
from io import BytesIO
from unittest.mock import patch

import pandas as pd
import pytest
from openpyxl import load_workbook

from application.tax_service import TaxCalculationService
from application.result_state import input_fingerprint, invalidate_result
from domain.models.broker_import import BrokerImportResult, WithholdingRecord
from domain.models.calculation_issue import CalculationIssue
from domain.models.deposit_record import DepositRecord
from domain.models.dividend_received import DividendReceivedRecord
from domain.models.position_record import PositionRecord
from domain.models.tax_summary import TaxSummary
from domain.models.exceptions import InventoryException, ParseException, TaxAssistantError
from domain.services.engine import FIFOEngine
from domain.services.tax_engine import TaxEngine
from domain.services.reporting import dividend_rows, capital_rows
from infrastructure.parsers.trade_parser import FutuTradeParser, _to_decimal
from infrastructure.exporters.excel_exporter import ExcelExporter
from tests.unit.test_fifo_engine import _trade
from tests.unit.test_tax_engine import _div, _match


def t(row, day, side, qty=10, price=100, **kwargs):
    return replace(_trade(row, date(2021, 1, day), "XYZ", side, qty, price), **kwargs)


def position(**kwargs):
    p = PositionRecord("期初", date(2020, 12, 31), "证券", "test", "test", "XYZ", "US", "USD", D(10), D(999), D(9990))
    return replace(p, **kwargs)


def data(**kwargs):
    d = BrokerImportResult(dividends=[_div("test", 100, 0, "CNY")], trade_file_present=True, report_years=[2021])
    return replace(d, **kwargs)


def workbook(sheets, name="test.xlsx"):
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        for sheet, rows in sheets.items():
            pd.DataFrame(rows).to_excel(writer, sheet_name=sheet, index=False)
    buf.name = name
    buf.seek(0)
    return buf


def uploads():
    dividend = workbook({"股息": [{"账户名称": "test", "年份": 2021, "币种": "CNY", "全年股息": 100,
                                 "全年利息": 0, "全年其他收入": 0}]}, "2021股息.xlsx")
    trade = workbook({
        "账户信息": [{"年份": 2021}],
        "证券-交易流水": [
            {"成交时间": "20210101", "账户名称": "test", "账户号码": "001", "品类": "证券", "代码名称": "XYZ",
             "交易所/市场": "US", "方向": "卖出", "数量/面值": -10, "价格": 150, "币种": "CNY"}],
        "证券-资金进出": [{"日期": "20210101", "币种": "CNY", "备注": "Withholding Tax", "变动金额": -10}],
    }, "2021_trade.xlsx")
    return {"dividend": dividend, "trade": trade}


def test_futu_amount_mismatch_marks_report_incomplete():
    files = uploads()
    trade = workbook({
        "账户信息": [{"年份": 2021}],
        "证券-交易流水": [{"成交时间": "20210102", "账户名称": "test", "账户号码": "001",
                     "品类": "证券", "代码名称": "XYZ", "交易所/市场": "US",
                     "方向": "买入", "数量/面值": 1, "价格": 10, "币种": "USD",
                     "成交金额": -12, "变动金额": -12, "总费用": 0}],
        "证券-资金进出": [],
    }, "2021_trade.xlsx")
    result = TaxCalculationService().calculate(files["dividend"], trade, 2021)
    assert not result.is_complete
    assert any(issue.code == "TRADE_AMOUNT_MISMATCH" for issue in result.issues)


@pytest.mark.parametrize("identity", [{"account_no": "other"}, {"currency": "HKD"}, {"market": "HK"},
                                      {"broker": "IBKR"}, {"instrument_id": "different"}])
def test_inventory_cannot_cross_identity(identity):
    engine = FIFOEngine()
    matches = engine.calculate([t(1, 1, "BUY"), t(2, 2, "SELL", price=150, **identity)], allow_incomplete=True)
    assert matches == []
    assert [i.code for i in engine.issues] == ["MISSING_COST"]


def test_each_affected_sell_reported_and_other_account_unaffected():
    engine = FIFOEngine()
    rows = [t(1, 1, "SELL"), t(2, 2, "BUY"), t(3, 3, "SELL"),
            t(4, 2, "BUY", account_no="safe"), t(5, 3, "SELL", price=150, account_no="safe")]
    matches = engine.calculate(rows, allow_incomplete=True, exchange_rate_func=lambda y, c: D(1))
    assert [i.source_row for i in engine.issues] == [1, 3]
    assert len(matches) == 1 and matches[0].account_no == "safe"
    assert matches[0].gain_original == D(500)
    engine.calculate([])
    assert engine.issues == [] and engine.warnings == []


def test_partial_oversell_does_not_emit_a_partial_match():
    engine = FIFOEngine()
    matches = engine.calculate([t(1, 1, "BUY"), t(2, 2, "SELL", qty=11)], allow_incomplete=True)
    assert matches == [] and engine.issues[0].quantity == 11


def test_same_day_real_time_wins_over_input_row_and_buy_priority():
    early_sell = t(8, 1, "SELL", trade_time=datetime(2021, 1, 1, 9))
    late_buy = t(2, 1, "BUY", trade_time=datetime(2021, 1, 1, 15))
    with pytest.raises(InventoryException):
        FIFOEngine().calculate([late_buy, early_sell])


def test_date_only_preserves_source_order_and_warns():
    engine = FIFOEngine()
    assert engine.calculate([t(2, 1, "BUY"), t(1, 1, "SELL")], allow_incomplete=True) == []
    assert engine.warnings and engine.issues


def test_cross_year_history_consumed_but_only_target_sales_use_fx():
    rows = [replace(t(1, 1, "BUY", qty=20), trade_date=date(2020, 1, 1)),
            replace(t(2, 2, "SELL", price=120), trade_date=date(2020, 2, 1)),
            t(3, 2, "SELL", price=150), replace(t(4, 3, "SELL"), trade_date=date(2022, 1, 1))]
    calls = []
    def rate(year, currency):
        calls.append(year)
        return D(7)
    m = FIFOEngine().calculate(rows, tax_year=2021, exchange_rate_func=rate)
    assert calls == [2021]
    assert len(m) == 1 and m[0].gain_cny == 3500


def test_unverified_opening_market_price_never_used_as_cost():
    engine = FIFOEngine()
    m = engine.calculate([t(1, 2, "SELL")], [position()], allow_incomplete=True)
    assert not m
    assert [i.code for i in engine.issues] == ["UNKNOWN_OPENING_COST", "UNCERTAIN_INVENTORY"]


def test_verified_opening_cost_and_overlap_guard():
    p = position(cost_basis_price=D(50))
    m = FIFOEngine().calculate([t(1, 2, "SELL")], [p], exchange_rate_func=lambda y, c: D(1))
    assert m[0].gain_original == 500
    with pytest.raises(InventoryException, match="重叠"):
        FIFOEngine().calculate([replace(t(1, 1, "BUY"), trade_date=date(2020, 1, 1))], [p])


def test_year_filter_covers_income_tax_cash_and_future_transactions():
    d = data(dividends=[_div("test", 100, 0, "CNY"), replace(_div("test", 200, 0, "CNY"), year=2022)],
             withholding=[WithholdingRecord(date(2021, 1, 1), "CNY", D(10)), WithholdingRecord(date(2022, 1, 1), "CNY", D(99))],
             deposits=[DepositRecord(date(2022, 1, 1), "test", "test", "CNY", D(100))])
    r = TaxCalculationService().calculate_imported(d, 2021)
    assert r.export_bundle.tax_summary.dividend_income_cny == 100
    assert r.export_bundle.tax_summary.foreign_tax_credit == 10
    assert r.export_bundle.deposits == []
    with pytest.raises(TaxAssistantError, match="多个年度"):
        TaxCalculationService().calculate_imported(d)


def test_tax_engine_year_guard_and_negative_credit_guard():
    s = TaxEngine().calculate([replace(_div("test", 100, 0), year=2020), _div("test", 200, 0)],
                              [replace(_match("XYZ", 100, 100), sell_date=date(2020, 1, 1))],
                              2021, lambda y, c: D(1), {"USD": D(-4)})
    assert s.dividend_income_cny == 200 and s.capital_gain_cny == 0 and s.foreign_tax_credit == 0


def test_unknown_income_year_requires_explicit_year_or_filename():
    d = data(dividends=[replace(_div("test", 100, 0), year=0)])
    with pytest.raises(TaxAssistantError, match="指定"):
        TaxCalculationService().calculate_imported(d)
    r = TaxCalculationService().calculate_imported(d, 2021)
    assert r.export_bundle.dividend_details[0].year == 2021 and d.dividends[0].year == 0
    with pytest.raises(TaxAssistantError, match="不一致"):
        TaxCalculationService().calculate_imported(d, 2021, income_file_year=2022)


def test_wrong_trade_report_year_is_rejected():
    with pytest.raises(TaxAssistantError, match="交易报告年度"):
        TaxCalculationService().calculate_imported(data(report_years=[2022]), 2021)


def test_signed_withholding_refund_and_year_filter():
    f = workbook({"资金进出": [
        {"日期": "20210101", "备注": "Withholding Tax", "币种": "USD", "变动金额": -10},
        {"日期": "20210102", "备注": "withholding tax reversal", "币种": "USD", "变动金额": 4},
        {"日期": "20220102", "备注": "Withholding Tax", "币种": "USD", "变动金额": -30}]})
    assert FutuTradeParser().parse_capital_flows(f, 2021) == {"USD": D(6)}


@pytest.mark.parametrize("value", ["NaN", "Infinity", "-Infinity", float("nan")])
def test_non_finite_numbers_rejected(value):
    with pytest.raises(ParseException):
        _to_decimal(value, "数量")


def test_trade_parser_preserves_time_account_and_reads_file_once():
    row = {"成交时间": "2021-01-01 15:30:01", "账户号码": "001", "账户名称": "test", "交易所/市场": "US",
           "品类": "证券", "代码名称": "XYZ", "方向": "买入", "数量/面值": 1, "价格": 100, "币种": "USD"}
    f = workbook({"交易流水": [row], "账户信息": [{"年份": 2021}], "资金进出": [{"备注": "not tax"}]})
    parser = FutuTradeParser()
    with patch.object(parser, "_open_excel", wraps=parser._open_excel) as opened:
        imported = parser.parse_all(f)
    assert opened.call_count == 1
    assert imported.trades[0].trade_time == datetime(2021, 1, 1, 15, 30, 1)
    assert imported.trades[0].account_no == "001" and imported.report_years == [2021]


def test_workbook_never_totals_unconverted_currencies_and_includes_issues():
    deps = [DepositRecord(date(2021, 1, 1), "test", "test", c, D(100)) for c in ("USD", "HKD")]
    issues = [CalculationIssue("MISSING_COST", "missing opening cost", symbol="XYZ", quantity=D(10))]
    b = ExcelExporter().build_report(TaxSummary(), [], [], deps, issues=issues, tax_year=2021)
    wb = load_workbook(BytesIO(b))
    assert "待复核记录" in wb.sheetnames
    assert "不完整" in wb["税务汇总"].cell(2, 2).value
    rows = list(wb["入金汇总"].values)
    assert len(rows) == 3 and {r[1] for r in rows[1:]} == {"USD", "HKD"}
    assert all(r[-1] == 100 for r in rows[1:])


def test_grouping_keeps_account_and_currency_separate():
    divs = [DividendReceivedRecord(date(2021, 1, 1), "test", account, currency, D(100), "XYZ", "")
            for account, currency in [("A", "USD"), ("B", "USD"), ("A", "HKD")]]
    assert len(dividend_rows(divs)) == 3
    matches = [replace(_match("XYZ", 100, 700), account_no=a, currency=c)
               for a, c in [("A", "USD"), ("B", "USD"), ("A", "HKD")]]
    assert len(capital_rows(matches)) == 3


def test_fingerprint_invalidates_content_filename_year_and_missing_file():
    files = uploads()
    a, b = files["dividend"], files["trade"]
    fingerprint = input_fingerprint(a, b, "2021")
    state = {"result": "old", "result_fingerprint": fingerprint}
    invalidate_result(state, fingerprint)
    assert state["result"] == "old"
    assert input_fingerprint(a, b, "2022") != fingerprint
    b.name = "changed.xlsx"
    assert input_fingerprint(a, b, "2021") != fingerprint
    b.write(b"changed")
    assert input_fingerprint(a, b, "2021") != fingerprint
    invalidate_result(state, None)
    assert not state


def test_app_rerun_preserves_warnings_but_changed_input_removes_download():
    from streamlit.testing.v1 import AppTest
    files = uploads()
    with patch("streamlit.file_uploader", side_effect=lambda *a, **kw: files[kw["key"]]):
        at = AppTest.from_file("../../app.py", default_timeout=30).run()
        at.button[0].click().run()
        assert not at.exception
        assert len(at.metric) == 7 and len(at.get("download_button")) == 1
        assert any("结果不完整" in e.value for e in at.error)
        warnings = [w.value for w in at.warning]
        assert any("缺少时间" in w for w in warnings)
        at.run()
        assert [w.value for w in at.warning] == warnings
        files["trade"].name = "new.xlsx"
        at.run()
        assert not at.metric and not at.get("download_button")
        at.text_input[0].set_value("invalid").run()
        at.button[0].click().run()
        assert not at.metric and not at.get("download_button")
        assert any("税款年度必须" in e.value for e in at.error)


def test_app_without_uploads_does_not_show_preexisting_result():
    from streamlit.testing.v1 import AppTest
    at = AppTest.from_file("../../app.py", default_timeout=30)
    at.session_state["result"] = "obsolete"
    at.run()
    assert not at.exception and not at.metric and not at.get("download_button")


def test_timezone_offsets_order_by_instant_and_mixed_timezones_rejected():
    buy = t(1, 1, "BUY", trade_time=datetime.fromisoformat("2021-01-01T09:00:00+00:00"))
    sell = t(2, 1, "SELL", trade_time=datetime.fromisoformat("2021-01-01T10:00:00+02:00"))
    with pytest.raises(InventoryException):
        FIFOEngine().calculate([buy, sell])
    with pytest.raises(InventoryException, match="混用"):
        FIFOEngine().calculate([buy, replace(sell, trade_time=datetime(2021, 1, 1, 10))])


def test_negative_net_withholding_marks_report_incomplete():
    r = TaxCalculationService().calculate_imported(data(withholding=[WithholdingRecord(date(2021, 1, 1), "CNY", D(-5))]), 2021)
    assert not r.is_complete and any(i.code == "NET_TAX_REFUND" for i in r.issues)
    assert r.export_bundle.tax_summary.foreign_tax_credit == 0


def test_missing_amount_and_duplicate_sheet_are_not_silent_zero():
    f = workbook({"资金进出": [{"日期": "20210101", "备注": "Withholding Tax", "币种": "USD"}]})
    with pytest.raises(ParseException, match="金额缺失"):
        FutuTradeParser().parse_capital_flows(f)
    f = workbook({"资金进出1": [{"x": 1}], "资金进出2": [{"x": 2}]})
    with pytest.raises(ParseException, match="多个"):
        FutuTradeParser().parse_capital_flows(f)


def test_calculation_failure_clears_previously_successful_result():
    from streamlit.testing.v1 import AppTest
    files = uploads()
    with patch("streamlit.file_uploader", side_effect=lambda *a, **kw: files[kw["key"]]):
        at = AppTest.from_file("../../app.py", default_timeout=30).run()
        at.button[0].click().run()
        assert at.metric
        with patch.object(TaxCalculationService, "calculate", side_effect=TaxAssistantError("test failure")):
            at.button[0].click().run()
        assert not at.metric and not at.get("download_button")
        assert any("test failure" in e.value for e in at.error)


def test_unresolved_transfer_prevents_using_affected_cost():
    r = TaxCalculationService().calculate_imported(data(
        trades=[t(1, 1, "BUY"), t(2, 2, "SELL", price=150)],
        issues=[CalculationIssue("UNSUPPORTED_TRANSFER", "split pending", account="test", symbol="XYZ", currency="USD")]), 2021)
    assert not r.export_bundle.match_records
    assert any(i.code == "UNCERTAIN_INVENTORY" for i in r.issues)


def test_complete_synthetic_report_has_exact_amounts_and_income_details():
    r = TaxCalculationService().calculate_imported(data(
        trades=[t(1, 1, "BUY", currency="CNY"), t(2, 2, "SELL", price=150, currency="CNY")],
        withholding=[WithholdingRecord(date(2021, 1, 1), "CNY", D(10)),
                     WithholdingRecord(date(2021, 1, 2), "CNY", D(-4))]), 2021)
    assert r.data_complete and not r.reconciliation_complete
    assert r.export_bundle.tax_summary.capital_gain_cny == 500
    assert r.export_bundle.tax_summary.foreign_tax_credit == 6
    assert r.export_bundle.tax_summary.total_supplement_tax == 114
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True)
    assert "收入明细" in wb.sheetnames and "待复核记录" not in wb.sheetnames
    rates = list(wb["汇率底稿"].iter_rows(values_only=True))
    assert rates[0] == ("年度", "币种", "汇率（兑人民币）", "来源")
    assert (2021, "CNY", "1", "系统固定汇率 CNY=1") in rates
    summary = {row[0]: row[1] for row in wb["税务汇总"].iter_rows(min_row=2, values_only=True)}
    assert summary["测算补税（待对账／临时）"] == 114
    notes = dict(wb["计算说明"].values)
    assert notes["税务规则版本"] == "V1-2026-09"
    assert "尚未按所得项目及国家／地区核对" in notes["规则适用边界"]
    wb.close()


def test_report_records_configured_exchange_rate_source():
    r = TaxCalculationService().calculate_imported(data(dividends=[_div("test", 100, 0, "USD")]), 2021)
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True)
    rows = list(wb["汇率底稿"].iter_rows(min_row=2, values_only=True))
    assert rows[0][0:3] == (2021, "USD", "6.4515")
    assert rows[0][3].endswith("exchange_rate.json")
    wb.close()
