"""金额守恒、对账状态和计算上下文的回归验证。"""
from io import BytesIO
from decimal import Decimal as D
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
import pytest

from tests.unit.test_ibkr import trade, cash, xml_file, service
from application.result_state import bind_calculation_context, invalidate_result


def balances(currency="USD", beginning="0", ending="496", **extra):
    return ("CashReport", "CashReportCurrency", dict(currency=currency, levelOfDetail="Currency",
        fromDate="20250101", toDate="20251231", startingCash=beginning, endingCash=ending, **extra))


def activity(pnl="496", ending="496", **sell):
    a = trade("T2", "SELL", dateTime="20250103;100000", tradeDate="20250103", **sell)
    if pnl is not None:
        a["fifoPnlRealized"] = pnl
    return xml_file(trades=[trade(), a], cash_rows=[], positions=[], extra=[balances(ending=ending)])


def test_original_proceeds_and_buy_basis_are_conserved():
    f = activity(pnl="496.02", ending="496.02", proceeds="1500.02", netCash="1498.02")
    r = service().calculate_ibkr([f], opening_zero=True)
    assert sum(m.sell_revenue for m in r.export_bundle.match_records) == D("1500.02")
    assert r.pnl_reconciliations[0].difference == 0
    assert r.is_complete


def test_one_original_purchase_split_across_sales_keeps_all_basis():
    from tests.unit.test_fifo_engine import _trade
    from dataclasses import replace
    from datetime import date
    from domain.services.engine import FIFOEngine
    buy = replace(_trade(1, date(2025, 1, 1), "TEST", "BUY", 3, 1), trade_amount=D("3.01"))
    sells = [_trade(i + 2, date(2025, 1, i + 2), "TEST", "SELL", 1, 2) for i in range(3)]
    matches = FIFOEngine().calculate([buy, *sells], exchange_rate_func=lambda y, c: D(1))
    assert sum(m.buy_cost for m in matches) == D("3.01")
    assert sum(m.sell_revenue for m in matches) == D("6")


def test_original_buy_amount_reconciles_with_reported_pnl():
    # 买入总额的微小差异必须保留到收益和现金核对。
    f = xml_file(trades=[trade(proceeds="-1000.01", netCash="-1002.01"),
        trade("T2", "SELL", fifoPnlRealized="495.99", dateTime="20250103;100000", tradeDate="20250103")],
        cash_rows=[], positions=[], extra=[balances(ending="495.99")])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert sum(m.buy_cost for m in r.export_bundle.match_records) == D("1000.01")
    assert r.is_complete


def test_missing_pnl_is_calculable_but_not_fully_reconciled():
    r = service().calculate_ibkr([activity(pnl=None)], opening_zero=True)
    assert r.data_complete and not r.reconciliation_complete and not r.is_complete
    assert r.pnl_reconciliations[0].status == "缺少券商收益"


def test_missing_cash_report_is_not_zero_or_complete():
    r = service().calculate_ibkr([xml_file(trades=[], cash_rows=[], positions=[])], opening_zero=True)
    assert r.data_complete and not r.is_complete
    assert r.cash_reconciliations[0]["状态"] == "无法核对"
    assert r.cash_reconciliations[0]["报告期末现金"] == ""


def test_cash_mismatch_is_an_issue_even_when_pnl_matches():
    r = service().calculate_ibkr([activity(ending="500")], opening_zero=True)
    assert r.pnl_reconciliations[0].status == "一致"
    assert any(i.code == "CASH_BALANCE_MISMATCH" for i in r.issues)
    assert not r.is_complete


def test_cash_uses_trade_date_balance_and_keeps_settled_balance_separate():
    f = xml_file(trades=[], cash_rows=[], positions=[], extra=[balances(ending="0", endingSettledCash="123")])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.is_complete and r.cash_reconciliations[0]["差额"] == "0"


def test_fx_cash_has_two_currency_legs_and_separate_commission():
    fx = trade("FX1", assetCategory="CASH", symbol="USD.HKD", currency="HKD", quantity="100",
               tradePrice="7.8", proceeds="-780", netCash="-780", ibCommission="-2", ibCommissionCurrency="USD")
    f = xml_file(trades=[fx], cash_rows=[], positions=[], extra=[
        balances(currency="USD", ending="98"), balances(currency="HKD", beginning="1000", ending="220")])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert len(r.cash_reconciliations) == 2
    assert all(row["状态"] == "一致" for row in r.cash_reconciliations)
    # 现金对得上并不表示未支持资产的税额已经计算。
    assert not r.is_complete


def test_other_fee_changes_cash_but_is_not_automatically_tax_deductible():
    f = xml_file(trades=[], cash_rows=[cash("F1", "Other Fees", "-2")], positions=[],
                 extra=[balances(beginning="10", ending="8")])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.cash_reconciliations[0]["状态"] == "一致"
    assert any(i.code == "UNSUPPORTED_CASH" for i in r.issues)


def test_duplicate_reports_are_reconciled_independently():
    r = service().calculate_ibkr([activity(), activity()], opening_zero=True)
    assert r.export_bundle.tax_summary.capital_gain_cny == D("496")
    assert len(r.cash_reconciliations) == 2 and r.is_complete


def test_wrong_cash_period_cannot_be_called_consistent():
    f = xml_file(trades=[], cash_rows=[], positions=[], extra=[balances(ending="0")])
    root = ET.fromstring(f.getvalue())
    root.find(".//CashReportCurrency").set("fromDate", "20250201")
    f = BytesIO(ET.tostring(root)); f.name = "wrong-period.xml"
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.cash_reconciliations[0]["状态"] == "无法核对" and not r.is_complete


def test_future_or_current_year_remains_provisional():
    f = activity()
    root = ET.fromstring(f.getvalue())
    for e in root.iter():
        for k, value in list(e.attrib.items()):
            if k in {"fromDate", "toDate", "dateTime", "tradeDate"}:
                e.set(k, value.replace("2025", "2099"))
    f = BytesIO(ET.tostring(root)); f.name = "future.xml"
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.data_complete and r.reconciliation_complete and r.provisional and not r.is_complete


def test_snapshot_and_status_are_exported_without_credentials():
    r = service().calculate_ibkr([activity()], opening_zero=True)
    assert r.snapshot["sources"][0]["sha256"]
    assert r.snapshot["context"]["fingerprint"]
    assert r.snapshot["exchange_rates"][0]["汇率（兑人民币）"] == "1"
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True, data_only=True)
    assert {"计算快照", "现金余额对账"} <= set(wb.sheetnames)
    assert dict(wb["计算说明"].values)["对账状态"] == "持仓、卖出盈亏、原币现金已核对"
    assert "token" not in str(r.snapshot).lower()
    wb.close()


def test_config_or_code_change_invalidates_same_upload():
    first = bind_calculation_context("same-files", {"fingerprint": "version1"})
    state = {"result": object(), "result_fingerprint": first}
    invalidate_result(state, bind_calculation_context("same-files", {"fingerprint": "version2"}))
    assert "result" not in state
    assert bind_calculation_context(None, {"fingerprint": "version1"}) is None


def test_context_hash_detects_actual_configuration_and_code_changes(tmp_path, monkeypatch):
    import application.calculation_context as context
    monkeypatch.setattr(context, "ROOT", tmp_path)
    (tmp_path / "config").mkdir()
    rate_file = tmp_path / "config" / "exchange_rate.json"
    rate_file.write_text('{"2025":{"USD":1}}', encoding="utf-8")
    first = context.calculation_context()["fingerprint"]
    rate_file.write_text('{"2025":{"USD":2}}', encoding="utf-8")
    second = context.calculation_context()["fingerprint"]
    (tmp_path / "app.py").write_text("# new version\n", encoding="utf-8")
    assert len({first, second, context.calculation_context()["fingerprint"]}) == 3
