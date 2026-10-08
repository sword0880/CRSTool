"""Synthetic Activity Flex reports: parsing, accounting boundaries and UI."""
from io import BytesIO
from pathlib import Path
from dataclasses import replace
from decimal import Decimal as D
from types import SimpleNamespace
from unittest.mock import patch
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from infrastructure.adapters.ibkr_adapter import IbkrReportAdapter
from application.tax_service import TaxCalculationService
from application.result_state import ibkr_fingerprint
from domain.models.exceptions import ParseException, TaxAssistantError, InventoryException


FIXTURE = Path(__file__).parents[1] / "fixtures" / "ibkr_activity.xml"
# 原有夹具未包含完整现金余额资料，数据完整性与全部对账完成分别断言。
# 全部对账状态回归见 test_architecture_boundaries.py。


def trade(id="T1", side="BUY", **overrides):
    a = dict(accountId="UTEST001", assetCategory="STK", conid="111", symbol="DEMO", currency="USD",
             tradeID=id, levelOfDetail="EXECUTION", dateTime="20250102;100000", tradeDate="20250102",
             buySell=side, quantity="10" if side == "BUY" else "-10", tradePrice="100" if side == "BUY" else "150",
             ibCommission="-2", ibCommissionCurrency="USD", taxes="0", multiplier="1",
             proceeds="-1000" if side == "BUY" else "1500", netCash="-1002" if side == "BUY" else "1498",
             openCloseIndicator="O" if side == "BUY" else "C", transactionType="ExchTrade")
    a.update(overrides)
    return a


def cash(id, kind="Dividends", amount="100", **overrides):
    a = dict(accountId="UTEST001", transactionID=id, dateTime="20250201;000000", currency="USD", type=kind,
             amount=amount, symbol="DEMO", conid="111", levelOfDetail="DETAIL")
    a.update(overrides)
    return a


def position(**overrides):
    a = dict(accountId="UTEST001", assetCategory="STK", conid="111", symbol="DEMO", currency="USD",
             levelOfDetail="SUMMARY", position="0", side="Long", multiplier="1")
    a.update(overrides)
    return a


def xml_file(trades=None, cash_rows=None, positions=None, start="20250101", end="20251231", name="activity.xml", extra=None):
    root = ET.Element("FlexQueryResponse", queryName="SYNTHETIC TEST ONLY", type="AF")
    statements = ET.SubElement(root, "FlexStatements", count="1")
    statement = ET.SubElement(statements, "FlexStatement", accountId="UTEST001", fromDate=start, toDate=end)
    for section, tag, rows in (("Trades", "Trade", trades), ("CashTransactions", "CashTransaction", cash_rows),
                                ("OpenPositions", "OpenPosition", positions)):
        if rows is not None:
            container = ET.SubElement(statement, section)
            for a in rows:
                ET.SubElement(container, tag, a)
    for section, tag, a in extra or []:
        container = statement.find(section)
        if container is None:
            container = ET.SubElement(statement, section)
        ET.SubElement(container, tag, a)
    buf = BytesIO(ET.tostring(root, encoding="utf-8", xml_declaration=True))
    buf.name = name
    return buf


def example():
    return xml_file(trades=[trade(), trade("T2", "SELL", dateTime="20250103;110000", tradeDate="20250103")],
                    cash_rows=[cash("C1"), cash("C2", "Withholding Tax", "-10"), cash("C3", "Withholding Tax", "4"),
                               cash("C4", "Broker Interest Received", "20"), cash("C5", "Deposits/Withdrawals", "500"),
                               cash("C6", "Deposits/Withdrawals", "-50")], positions=[])


def service():
    class ConfirmedScopeService(TaxCalculationService):
        def calculate_ibkr(self, report_files, opening_file=None, tax_year=None, opening_zero=False,
                           source_scope_confirmed=True):
            return super().calculate_ibkr(report_files, opening_file, tax_year, opening_zero,
                                          source_scope_confirmed=source_scope_confirmed)

    return ConfirmedScopeService(exchange_rate_repo=SimpleNamespace(get_rate=lambda y, c: D(1)))


def opening_file(**overrides):
    a = position(levelOfDetail="LOT", position="5", costBasisMoney="400", openDateTime="20230101;100000",
                 originatingTransactionID="OPEN1")
    a.update(overrides)
    return xml_file(positions=[a, position(levelOfDetail="LOT", position="5", costBasisMoney="600",
                    openDateTime="20240101;100000", originatingTransactionID="OPEN2")],
                    start="20241231", end="20241231", name="opening.xml")


def test_complete_activity_exact_amounts_and_audit_sheets():
    r = service().calculate_ibkr([example()], opening_zero=True)
    assert r.data_complete
    s = r.export_bundle.tax_summary
    assert s.capital_gain_cny == 496 and s.dividend_income_cny == 100 and s.interest_income_cny == 20
    assert s.foreign_tax_credit == 6 and s.total_supplement_tax == D("117.20")
    assert len(r.export_bundle.deposits) == 1
    assert r.export_bundle.match_records[0].buy_record_id == "T1"
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True)
    assert {"导入来源", "现金事件明细", "资本利得明细"} <= set(wb.sheetnames)
    assert wb["现金事件明细"].max_row == 7
    assert len(wb["导入来源"].cell(2, 2).value) == 64
    wb.close()


def test_checked_in_fixture_matches_builder_and_no_income_is_allowed():
    assert FIXTURE.read_bytes() == example().getvalue()
    r = service().calculate_ibkr([xml_file(trades=[], cash_rows=[], positions=[])], opening_zero=True)
    assert r.data_complete and r.export_bundle.tax_summary.total_supplement_tax == 0


def test_ibkr_scope_requires_explicit_confirmation_and_is_recorded_in_report():
    unconfirmed = service().calculate_ibkr([example()], opening_zero=True, source_scope_confirmed=False)
    assert not unconfirmed.data_complete
    assert any(i.code == "SOURCE_SCOPE_UNCONFIRMED" for i in unconfirmed.issues)
    wb = load_workbook(BytesIO(unconfirmed.report_bytes), read_only=True)
    notes = dict(wb["计算说明"].values)
    assert notes["IBKR 来源范围确认"] == "未确认导出范围"
    wb.close()
    confirmed = service().calculate_ibkr([example()], opening_zero=True)
    wb = load_workbook(BytesIO(confirmed.report_bytes), read_only=True)
    assert dict(wb["计算说明"].values)["IBKR 来源范围确认"] == "用户已确认导出范围"
    wb.close()


def test_missing_cash_section_is_not_zero_income():
    with pytest.raises(TaxAssistantError, match="缺少收入"):
        service().calculate_ibkr([xml_file(trades=[], positions=[])], opening_zero=True)


def test_reports_deduplicate_by_account_and_record_id():
    f = example()
    r = service().calculate_ibkr([f, f], opening_zero=True)
    assert r.data_complete and r.export_bundle.tax_summary.capital_gain_cny == 496
    assert any("重复" in w for w in r.warnings)


def test_conflicting_duplicate_trade_or_cash_fails():
    with pytest.raises(ParseException, match="冲突"):
        IbkrReportAdapter().parse_files([example(), xml_file(trades=[trade(tradePrice="101")], cash_rows=[])])
    with pytest.raises(ParseException, match="冲突"):
        IbkrReportAdapter().parse_files([example(), xml_file(trades=[], cash_rows=[cash("C1", amount="101")])])


def test_summary_and_closed_lot_not_counted_as_executions():
    f = xml_file(trades=[trade(), trade("T2", "SELL", dateTime="20250103;110000", tradeDate="20250103"),
                        trade("T3", levelOfDetail="CLOSED_LOT"), trade("T4", levelOfDetail="ORDER")], cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.data_complete and r.export_bundle.tax_summary.capital_gain_cny == 496
    f = xml_file(trades=[trade(levelOfDetail="ORDER")], cash_rows=[], positions=[])
    assert any(i.code == "MISSING_EXECUTIONS" for i in service().calculate_ibkr([f], opening_zero=True).issues)


@pytest.mark.parametrize("changes", [{"assetCategory": "OPT"}, {"origTradeID": "T0"}, {"ibCommissionCurrency": "HKD"},
                                    {"multiplier": "100"}, {"ibCommission": "1"}, {"taxes": "-1"}, {"notes": "Ca"}])
def test_unsupported_trade_blocks_affected_instrument(changes):
    f = xml_file(trades=[trade(**changes), trade("T2", "SELL", dateTime="20250103;110000", tradeDate="20250103")], cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert not r.data_complete and not r.export_bundle.match_records


def test_opening_lots_use_original_dates_and_total_cost_not_market_value():
    sell = trade("T2", "SELL")
    f = xml_file(trades=[sell], cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_file=opening_file(markPrice="9999"))
    assert r.data_complete
    assert r.export_bundle.tax_summary.capital_gain_cny == 498
    assert [m.buy_cost for m in r.export_bundle.match_records] == [400, 600]
    assert [m.buy_date.year for m in r.export_bundle.match_records] == [2023, 2024]
    assert r.export_bundle.match_records[0].buy_source_file == "opening.xml"


def test_no_opening_source_does_not_claim_complete():
    r = service().calculate_ibkr([example()])
    assert not r.data_complete and any(i.code == "OPENING_UNCONFIRMED" for i in r.issues)


def test_opening_snapshot_requires_lots_correct_year_and_nonoverlap():
    with pytest.raises(ParseException, match="LOT"):
        service().calculate_ibkr([example()], xml_file(positions=[position()], start="20241231", end="20241231"))
    with pytest.raises(TaxAssistantError, match="截止日期"):
        service().calculate_ibkr([example()], xml_file(positions=[], start="20250101", end="20251231"))
    with pytest.raises(ParseException, match="同时声明"):
        service().calculate_ibkr([example()], opening_file(), opening_zero=True)
    history = trade(dateTime="20240101;100000", tradeDate="20240101")
    with pytest.raises(InventoryException, match="重叠"):
        service().calculate_ibkr([xml_file(trades=[history], cash_rows=[], positions=[], start="20240101")], opening_file(), tax_year=2025)


def test_monthly_gaps_and_final_quantity_mismatch_mark_incomplete():
    f = xml_file(trades=[], cash_rows=[], positions=[], start="20250201")
    r = service().calculate_ibkr([f], opening_zero=True)
    assert any(i.code == "INCOMPLETE_PERIOD" for i in r.issues)
    r = service().calculate_ibkr([xml_file(trades=[trade()], cash_rows=[], positions=[position(position="9")])], opening_zero=True)
    assert any(i.code == "POSITION_MISMATCH" for i in r.issues)


def test_contiguous_reports_cover_year_without_double_counting():
    f1 = xml_file(trades=[trade()], cash_rows=[], positions=[position(position="10")], end="20250630", name="H1.xml")
    f2 = xml_file(trades=[trade("T2", "SELL", dateTime="20250701;100000", tradeDate="20250701")], cash_rows=[],
                  positions=[], start="20250701", name="H2.xml")
    r = service().calculate_ibkr([f2, f1], opening_zero=True)
    assert r.data_complete and r.export_bundle.tax_summary.capital_gain_cny == 496


def test_multi_account_same_ids_do_not_merge_inventory_or_income():
    a = example()
    b = BytesIO(a.getvalue().replace(b"UTEST001", b"UTEST002")); b.name = "account2.xml"
    r = service().calculate_ibkr([a, b], opening_zero=True)
    assert r.data_complete and r.export_bundle.tax_summary.capital_gain_cny == 992
    assert len(r.export_bundle.match_records) == 2


def test_no_cross_account_cost_and_event_account_must_match():
    root = ET.fromstring(example().getvalue())
    root.find(".//Trade").set("accountId", "UTEST002")
    f = BytesIO(ET.tostring(root)); f.name = "bad.xml"
    with pytest.raises(ParseException, match="账户不一致"):
        IbkrReportAdapter().parse_files([f])


def test_dividend_accrual_not_counted_and_unknown_cash_retained():
    f = xml_file(trades=[], cash_rows=[cash("C1", "Payment In Lieu Of Dividends", "100")], positions=[],
                 extra=[("ChangeInDividendAccruals", "ChangeInDividendAccrual", {"grossAmount": "500"})])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.export_bundle.tax_summary.dividend_income_cny == 0
    assert any(i.code == "UNSUPPORTED_CASH" for i in r.issues)
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True)
    assert wb["现金事件明细"].max_row == 2
    wb.close()


def test_negative_interest_expense_not_net_against_received_interest():
    f = xml_file(trades=[], cash_rows=[cash("I1", "Broker Interest Received", "20"), cash("I2", "Broker Interest Paid", "-50")], positions=[])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.export_bundle.tax_summary.interest_income_cny == 20 and not r.data_complete


def test_date_only_cash_supported_and_reversal_reduces_income():
    f = xml_file(trades=[], cash_rows=[cash("C1", dateTime="20250201"), cash("C2", amount="-10")], positions=[])
    assert service().calculate_ibkr([f], opening_zero=True).export_bundle.tax_summary.dividend_income_cny == 90


def test_cash_report_original_currency_is_reconciled_to_detail():
    summary = {"accountId": "UTEST001", "levelOfDetail": "Currency", "currency": "USD",
               "fromDate": "20250101", "toDate": "20251231", "dividends": "100",
               "brokerInterest": "0", "withholdingTax": "0"}
    f = xml_file(trades=[], cash_rows=[cash("C1")], positions=[],
                 extra=[("CashReport", "CashReportCurrency", summary)])
    assert not any(i.code == "CASH_REPORT_MISMATCH" for i in IbkrReportAdapter().parse_files([f]).issues)
    f = xml_file(trades=[], cash_rows=[cash("C1")], positions=[],
                 extra=[("CashReport", "CashReportCurrency", {**summary, "dividends": "99"})])
    result = service().calculate_ibkr([f], opening_zero=True)
    assert any(i.code == "CASH_REPORT_MISMATCH" for i in result.issues)
    assert not result.data_complete


def test_cash_report_missing_field_or_currency_is_not_silent():
    row = {"accountId": "UTEST001", "levelOfDetail": "Currency", "currency": "USD",
           "fromDate": "20250101", "toDate": "20251231"}
    f = xml_file(trades=[], cash_rows=[cash("C1")], positions=[],
                 extra=[("CashReport", "CashReportCurrency", row)])
    assert any(i.code == "CASH_REPORT_MISSING_FIELD" for i in IbkrReportAdapter().parse_files([f]).issues)
    f = xml_file(trades=[], cash_rows=[cash("C1")], positions=[],
                 extra=[("CashReport", "CashReportCurrency", {**row, "currency": "HKD"})])
    assert any(i.code == "CASH_REPORT_MISSING_CURRENCY" for i in IbkrReportAdapter().parse_files([f]).issues)


def test_prior_year_unsupported_fx_trade_does_not_mark_target_year_incomplete():
    fx = trade("FX1", "SELL", assetCategory="CASH", conid="FX", currency="HKD",
               ibCommissionCurrency="HKD", dateTime="20241231;100000", tradeDate="20241231")
    f = xml_file(trades=[fx, trade(), trade("T2", "SELL", dateTime="20250103;100000",
                                          tradeDate="20250103")], cash_rows=[], positions=[], start="20241231")
    imported = IbkrReportAdapter().parse_files([f], opening_zero=True)
    issue = next(i for i in imported.issues if i.code == "UNSUPPORTED_ASSET")
    assert issue.date.isoformat() == "2024-12-31"
    result = service().calculate_imported(imported, 2025)
    assert not any(i.code == "UNSUPPORTED_ASSET" for i in result.issues)


def test_prior_year_missing_section_does_not_leak_into_target_year():
    prior = xml_file(trades=[], positions=[], start="20240101", end="20241231", name="same.xml")
    target = xml_file(trades=[trade(), trade("T2", "SELL", dateTime="20250103;100000",
                                                  tradeDate="20250103")],
                      cash_rows=[], positions=[], name="same.xml")
    result = service().calculate_ibkr([prior, target], tax_year=2025, opening_zero=True)
    assert not any(i.code == "MISSING_SECTION" for i in result.issues)


@pytest.mark.parametrize("field", ["conid", "tradeID", "ibCommission", "taxes", "multiplier"])
def test_required_trade_fields_fail_with_actionable_message(field):
    a = trade(); del a[field]
    with pytest.raises(ParseException, match=field):
        IbkrReportAdapter().parse_files([xml_file(trades=[a], cash_rows=[], positions=[])])


@pytest.mark.parametrize("value", ["NaN", "Infinity", "not-a-number"])
def test_nonfinite_amount_rejected(value):
    with pytest.raises(ParseException):
        IbkrReportAdapter().parse_files([xml_file(trades=[], cash_rows=[cash("C1", amount=value)])])


def test_xml_errors_receipts_dtd_and_utf16_entity_rejected():
    for payload in (b"not XML", b"<FlexStatementResponse><Status>Success</Status></FlexStatementResponse>",
                    b'<!DOCTYPE x [<!ENTITY secret "expanded">]><FlexQueryResponse>&secret;</FlexQueryResponse>',
                    '<!DOCTYPE x [<!ENTITY secret "expanded">]><FlexQueryResponse>&secret;</FlexQueryResponse>'.encode("utf-16")):
        f = BytesIO(payload); f.name = "bad.xml"
        with pytest.raises(ParseException):
            IbkrReportAdapter().parse_files([f])


def test_input_fingerprint_includes_opening_and_confirmation():
    f = example()
    a = ibkr_fingerprint([f], None, "2025", False)
    assert a != ibkr_fingerprint([f], None, "2025", True)
    assert a != ibkr_fingerprint([f], opening_file(), "2025", False)
    assert a != ibkr_fingerprint([f], None, "2024", False)
    assert a != ibkr_fingerprint([f], None, "2025", False, True)


def test_streamlit_ibkr_upload_calculate_and_broker_switch():
    from streamlit.testing.v1 import AppTest
    files = {"dividend": None, "trade": None, "ibkr_reports": [example()], "ibkr_opening": None}
    with patch("streamlit.file_uploader", side_effect=lambda *a, **kw: files[kw["key"]]):
        at = AppTest.from_file("../../app.py", default_timeout=30).run()
        at.selectbox[0].select("IBKR").run()
        at.checkbox[0].check().run()
        at.checkbox[1].check().run()
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert not at.exception and len(at.metric) == 7
        assert len(at.get("download_button")) == 1
        assert not at.error
        at.checkbox[0].uncheck().run()
        assert not at.metric and not at.get("download_button")
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert any("不完整" in e.value for e in at.error)
        at.selectbox[0].select("富途").run()
        assert not at.exception and not at.metric and not at.get("download_button")


def test_ibkr_ui_opening_zero_ignores_uploaded_summary():
    from streamlit.testing.v1 import AppTest

    files = {"dividend": None, "trade": None, "ibkr_reports": [example()],
             "ibkr_opening": example()}
    with patch("streamlit.file_uploader", side_effect=lambda *a, **kw: files[kw["key"]]):
        at = AppTest.from_file("../../app.py", default_timeout=30).run()
        at.selectbox[0].select("IBKR").run()
        at.checkbox[0].check().run()
        at.checkbox[1].check().run()
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert not at.exception
        assert not any("期初成本需 LOT" in error.value for error in at.error)


def test_semantically_identical_decimal_values_deduplicate():
    f = xml_file(trades=[trade(tradePrice="100.000", quantity="10.0")], cash_rows=[], positions=[])
    r = service().calculate_ibkr([example(), f], opening_zero=True)
    assert r.data_complete and r.export_bundle.tax_summary.capital_gain_cny == 496


def test_cash_summary_missing_id_and_report_date_are_checked():
    r = service().calculate_ibkr([xml_file(trades=[], cash_rows=[cash("C1", levelOfDetail="SUMMARY")], positions=[])], opening_zero=True)
    assert any(i.code == "UNKNOWN_CASH_LEVEL" for i in r.issues)
    a = cash("C1"); del a["transactionID"]
    with pytest.raises(ParseException, match="transactionID"):
        IbkrReportAdapter().parse_files([xml_file(trades=[], cash_rows=[a])])
    with pytest.raises(ParseException, match="日期"):
        IbkrReportAdapter().parse_files([xml_file(trades=[], cash_rows=[], positions=[position(reportDate="20241231")])])


def test_snapshot_conflict_and_opening_lot_summary_mismatch():
    with pytest.raises(ParseException, match="快照冲突"):
        IbkrReportAdapter().parse_files([example(), xml_file(trades=[], cash_rows=[], positions=[position(position="2")])])
    f = opening_file()
    root = ET.fromstring(f.getvalue())
    ET.SubElement(root.find(".//OpenPositions"), "OpenPosition", position(position="11", costBasisMoney="1000"))
    f = BytesIO(ET.tostring(root)); f.name = "opening.xml"
    with pytest.raises(ParseException, match="SUMMARY"):
        service().calculate_ibkr([example()], f)


def test_unimplemented_fx_and_corporate_action_are_visible():
    f = xml_file(trades=[], cash_rows=[], positions=[], extra=[("FxTransactions", "FxTransaction", {"realizedPL": "123"})])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert not r.data_complete and any(i.code == "UNSUPPORTED_SECTION" for i in r.issues)
    f = xml_file(trades=[trade(), trade("T2", "SELL", dateTime="20250103;110000", tradeDate="20250103")],
                 cash_rows=[], positions=[], extra=[("CorporateActions", "CorporateAction", {"accountId": "UTEST001", "conid": "111", "symbol": "RENAMED"})])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert not r.export_bundle.match_records and not r.data_complete


def test_ibkr_fx_rate_to_base_is_not_used_as_cny_rate():
    f = xml_file(trades=[], cash_rows=[cash("C1", fxRateToBase="123")], positions=[])
    assert service().calculate_ibkr([f], opening_zero=True).export_bundle.tax_summary.dividend_income_cny == 100


def test_file_size_limit_and_namespace(monkeypatch):
    import infrastructure.adapters.ibkr_adapter as module
    with monkeypatch.context() as m:
        m.setattr(module, "MAX_FILE_BYTES", 10)
        with pytest.raises(ParseException, match="25 MB"):
            IbkrReportAdapter().parse_files([example()])
    f = BytesIO(example().getvalue().replace(b'<FlexQueryResponse ', b'<FlexQueryResponse xmlns="urn:test" ')); f.name = "namespaced.xml"
    assert service().calculate_ibkr([f], opening_zero=True).data_complete


def test_unsupported_model_and_wrong_record_period():
    with pytest.raises(ParseException, match="Model"):
        IbkrReportAdapter().parse_files([xml_file(trades=[trade(model="A")], cash_rows=[])])
    with pytest.raises(ParseException, match="超出"):
        IbkrReportAdapter().parse_files([xml_file(trades=[trade(dateTime="20240102;100000", tradeDate="20240102")], cash_rows=[])])


def test_ibkr_ui_opening_file_and_error_clear_previous_result():
    from streamlit.testing.v1 import AppTest
    files = {"dividend": None, "trade": None, "ibkr_reports": [xml_file(trades=[trade("T2", "SELL")], cash_rows=[], positions=[])], "ibkr_opening": opening_file()}
    with patch("streamlit.file_uploader", side_effect=lambda *a, **kw: files[kw["key"]]):
        at = AppTest.from_file("../../app.py", default_timeout=30).run()
        at.selectbox[0].select("IBKR").run()
        assert not at.checkbox[0].disabled
        at.checkbox[1].check().run()
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert not at.exception and not at.error and at.metric
        files["ibkr_reports"] = [BytesIO(b"bad")]
        files["ibkr_reports"][0].name = "bad.xml"
        at.run()
        assert not at.metric
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert at.error and not at.get("download_button")
