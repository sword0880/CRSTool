"""Execution-level IBKR P/L checks use synthetic, independent broker values."""
from dataclasses import replace
from decimal import Decimal as D
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

from tests.unit.test_ibkr import trade, xml_file, example, opening_file, service
from infrastructure.adapters.ibkr_adapter import IbkrReportAdapter
from domain.models.exceptions import ParseException
from application.ibkr_review import reconcile_realized_pnl


def pnl_file(amount="496", **overrides):
    sell = trade("T2", "SELL", dateTime="20250103;110000", tradeDate="20250103", **overrides)
    if amount is not None:
        sell["fifoPnlRealized"] = amount
    return xml_file(trades=[trade(), sell], cash_rows=[], positions=[])


def test_exact_pnl_and_duplicate_reports_preserve_one_comparison():
    f = pnl_file()
    r = service().calculate_ibkr([f, f], opening_zero=True)
    assert r.data_complete and len(r.pnl_reconciliations) == 1
    row = r.pnl_reconciliations[0]
    assert (row.calculated, row.reported.amount, row.difference, row.status) == (496, 496, 0, "一致")
    assert r.export_bundle.tax_summary.capital_gain_cny == 496
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True)
    values = list(wb["已实现盈亏对账"].values)
    record = dict(zip(values[0], values[1]))
    assert record["成交ID"] == "T2" and record["差额（本系统减券商）"] == 0
    assert record["来源文件"] == "activity.xml" and record["记录序号"] == 2
    wb.close()


@pytest.mark.parametrize("amount,expected", [("495.98", "一致"), ("496.02", "一致"), ("495.979", "存在差异"), ("496.021", "存在差异"), ("0", "存在差异"), ("-10", "存在差异")])
def test_signed_tolerance_and_broker_amount_never_overrides_calculation(amount, expected):
    r = service().calculate_ibkr([pnl_file(amount)], opening_zero=True)
    row = r.pnl_reconciliations[0]
    assert row.status == expected and row.difference == D(496) - D(amount)
    assert r.export_bundle.tax_summary.capital_gain_cny == 496
    assert r.data_complete == (expected == "一致")
    if expected != "一致":
        issue = next(i for i in r.issues if i.code == "REALIZED_PNL_MISMATCH")
        assert issue.record_id == "T2" and issue.source_file == "activity.xml"
        assert issue.instrument_id == "111" and issue.currency == "USD"


@pytest.mark.parametrize("value", [None, "", "   "])
def test_missing_pnl_remains_blank_and_is_not_zero(value):
    r = service().calculate_ibkr([pnl_file(value)], opening_zero=True)
    row = r.pnl_reconciliations[0]
    assert row.status == "缺少券商收益" and row.reported.amount is None and row.difference is None
    assert row.calculated == 496 and r.data_complete
    assert any("1 笔卖出未提供" in w for w in r.warnings)
    wb = load_workbook(BytesIO(r.report_bytes), read_only=True)
    values = list(wb["已实现盈亏对账"].values)
    record = dict(zip(values[0], values[1]))
    assert record["券商收益（原币）"] is None and record["差额（本系统减券商）"] is None
    wb.close()


@pytest.mark.parametrize("value", ["NaN", "Infinity", "not-money"])
def test_invalid_reported_pnl_rejected(value):
    with pytest.raises(ParseException):
        service().calculate_ibkr([pnl_file(value)], opening_zero=True)


def test_changed_pnl_for_same_id_is_a_conflict_not_a_duplicate():
    with pytest.raises(ParseException, match="冲突"):
        IbkrReportAdapter().parse_files([pnl_file("496"), pnl_file("497")])
    data = IbkrReportAdapter().parse_files([pnl_file("496"), pnl_file("496.000")])
    assert len(data.realized_pnl) == 1


def test_sale_matching_multiple_opening_lots_compared_once():
    f = xml_file(trades=[trade("T2", "SELL", fifoPnlRealized="498")], cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_file())
    assert len(r.export_bundle.match_records) == 2 and len(r.pnl_reconciliations) == 1
    assert r.pnl_reconciliations[0].calculated == 498 and r.pnl_reconciliations[0].status == "一致"


@pytest.mark.parametrize("missing_cost", [True, False])
def test_uncalculated_or_unsupported_sale_does_not_claim_zero_or_match(missing_cost):
    sells = [trade("T2", "SELL", fifoPnlRealized="0")]
    if not missing_cost:
        sells[0]["assetCategory"] = "OPT"
    r = service().calculate_ibkr([xml_file(trades=sells, cash_rows=[], positions=[])], opening_zero=True)
    row = r.pnl_reconciliations[0]
    assert row.status == ("无法核对" if missing_cost else "未支持资产")
    assert row.calculated is None and row.difference is None
    assert row.reported.amount == 0 and not r.data_complete


def test_fx_sell_is_separate_from_stock_reconciliation_denominator():
    fx = trade("FX1", "SELL", assetCategory="CASH", conid="FX", currency="HKD",
               ibCommissionCurrency="HKD", openCloseIndicator="",
               dateTime="20250104;100000", tradeDate="20250104", fifoPnlRealized="0")
    f = xml_file(trades=[trade(), trade("T2", "SELL", fifoPnlRealized="496",
                                      dateTime="20250103;110000", tradeDate="20250103"), fx],
                 cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert [row.status for row in r.pnl_reconciliations] == ["一致", "未支持资产"]
    assert sum(row.status == "一致" for row in r.pnl_reconciliations) == 1
    assert sum(row.status != "未支持资产" for row in r.pnl_reconciliations) == 1


def test_two_offsetting_differences_do_not_disappear_in_annual_net():
    f = xml_file(trades=[trade(), trade("T2", "SELL", fifoPnlRealized="500", dateTime="20250103;100000", tradeDate="20250103"),
                        trade("T3", dateTime="20250104;100000", tradeDate="20250104"),
                        trade("T4", "SELL", fifoPnlRealized="492", dateTime="20250105;100000", tradeDate="20250105")], cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert [row.difference for row in r.pnl_reconciliations] == [-4, 4]
    assert len([i for i in r.issues if i.code == "REALIZED_PNL_MISMATCH"]) == 2
    assert not r.data_complete and r.export_bundle.tax_summary.capital_gain_cny == 992


def test_same_id_in_other_account_or_currency_cannot_cross_match():
    root = ET.fromstring(pnl_file().getvalue())
    other = ET.fromstring(pnl_file("490").getvalue()).find(".//FlexStatement")
    other.set("accountId", "UTEST002")
    for elem in other.findall(".//Trade"):
        elem.set("accountId", "UTEST002")
        elem.set("currency", "HKD")
        elem.set("ibCommissionCurrency", "HKD")
    root.find("FlexStatements").append(other)
    f = BytesIO(ET.tostring(root)); f.name = "multi.xml"
    r = service().calculate_ibkr([f], opening_zero=True)
    assert [(row.reported.account, row.reported.currency, row.difference) for row in r.pnl_reconciliations] == [
        ("UTEST001", "USD", D(0)), ("UTEST002", "HKD", D(6))]


def test_reconciliation_requires_complete_matched_quantity_and_target_year():
    svc = service()
    data = IbkrReportAdapter().parse_files([pnl_file()], opening_zero=True)
    result = svc.calculate_imported(data, 2025)
    partial = [replace(result.export_bundle.match_records[0], sell_quantity=D(5))]
    rows, _, _ = reconcile_realized_pnl(data, 2025, svc.fifo_engine, partial)
    assert rows[0].status == "无法核对" and rows[0].calculated is None
    rows, _, _ = reconcile_realized_pnl(data, 2024, svc.fifo_engine, result.export_bundle.match_records)
    assert not rows


def test_sell_with_zero_realized_gain_is_a_valid_reference():
    f = xml_file(trades=[trade(), trade("T2", "SELL", tradePrice="100.4", proceeds="1004", netCash="1002",
                 dateTime="20250103;100000", tradeDate="20250103", fifoPnlRealized="0")], cash_rows=[], positions=[])
    r = service().calculate_ibkr([f], opening_zero=True)
    assert r.data_complete and r.pnl_reconciliations[0].status == "一致"
    assert r.pnl_reconciliations[0].calculated == 0


def test_ui_reconciliation_and_changed_report_clear_old_comparison():
    from streamlit.testing.v1 import AppTest
    files = {"dividend": None, "trade": None, "ibkr_reports": [pnl_file("500")], "ibkr_opening": None}
    with patch("streamlit.file_uploader", side_effect=lambda *a, **kw: files[kw["key"]]):
        at = AppTest.from_file("../../app.py", default_timeout=30).run()
        at.selectbox[0].select("IBKR").run()
        at.checkbox[0].check().run()
        at.checkbox[1].check().run()
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert not at.exception and any("不完整" in e.value for e in at.error)
        assert any("已实现盈亏对账（0/1 笔一致）" in e.label for e in at.expander)
        frame = next(d.value for d in at.dataframe if "券商收益（原币）" in d.value.columns)
        assert frame.iloc[0]["差额（本系统减券商）"] == -4
        files["ibkr_reports"] = [pnl_file()]
        at.run()
        assert not at.metric and not any("已实现盈亏对账" in e.label for e in at.expander)
        next(button for button in at.button if button.label == "开始计算").click().run()
        assert not at.error and any("已实现盈亏对账（1/1 笔一致）" in e.label for e in at.expander)


def test_reconciliation_fixture_matches_calculation():
    fixture = Path(__file__).parents[1] / "fixtures" / "ibkr_reconciliation.xml"
    assert fixture.read_bytes() == pnl_file().getvalue()
    assert service().calculate_ibkr([fixture], opening_zero=True).pnl_reconciliations[0].status == "一致"
